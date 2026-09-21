from __future__ import annotations

import io
import unittest

from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from docgen.anonymizer.verified_engine import build_verified_rules_engine
from docgen.document_formats import detect_document_format, validate_document
from docgen.renderer import render_document
from docgen.template_service import analyze_template


def _pdf_bytes(text: str) -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output)
    document.drawString(72, 720, text)
    document.save()
    return output.getvalue()


def _pptx_bytes(text: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    shape.text = text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _scanned_pdf_bytes(text: str) -> bytes:
    image = Image.new("RGB", (1200, 500), "white")
    drawing = ImageDraw.Draw(image)
    drawing.text(
        (80, 180),
        text,
        fill="black",
        font=ImageFont.truetype("DejaVuSans.ttf", 48),
    )
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=(600, 250))
    document.drawImage(ImageReader(image), 0, 0, width=600, height=250)
    document.save()
    return output.getvalue()


class DocumentFormatTests(unittest.TestCase):
    def test_format_is_detected_from_content_and_must_match_extension(self) -> None:
        pdf = _pdf_bytes("Documento")
        self.assertEqual(detect_document_format(pdf), "pdf")
        self.assertEqual(
            validate_document("documento.pdf", pdf, max_expanded_bytes=1_000_000)[1],
            "pdf",
        )
        with self.assertRaisesRegex(ValueError, "extensión"):
            validate_document("documento.docx", pdf, max_expanded_bytes=1_000_000)


class PdfIntegrationTests(unittest.TestCase):
    def test_scanned_pdf_uses_the_local_ocr_route(self) -> None:
        source = _scanned_pdf_bytes("DNI 12345678Z - ana@example.com")
        analysis = build_verified_rules_engine().analyze(source)

        self.assertFalse(analysis.coverage_issues)
        self.assertTrue(
            any(block.location.part.endswith("/ocr") for block in analysis.blocks)
        )
        self.assertIn("12345678Z", {finding.value for finding in analysis.findings})

    def test_pdf_can_be_anonymized_and_rebuilt_without_a_text_layer(self) -> None:
        source = _pdf_bytes("DNI 12345678Z y email ana@example.com")
        engine = build_verified_rules_engine()
        analysis = engine.analyze(source)

        output, applied = engine.apply(
            source,
            analysis,
            selected_group_ids={group.id for group in analysis.groups},
            replacements={},
            remove_unsupported_content=False,
        )

        self.assertEqual({item.entity_type for item in applied}, {"ES_DNI", "EMAIL"})
        self.assertEqual(len(PdfReader(io.BytesIO(output)).pages), 1)
        self.assertEqual(
            "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(output)).pages),
            "",
        )

    def test_pdf_template_variables_are_discovered_and_rendered(self) -> None:
        source = _pdf_bytes("Cliente: {{ cliente }}")
        analysis = analyze_template(source)
        output = render_document(source, {"cliente": "Acme"})

        self.assertEqual(analysis.variables, ("cliente",))
        self.assertEqual(analysis.document_format, "pdf")
        self.assertEqual(detect_document_format(output), "pdf")


class PptxIntegrationTests(unittest.TestCase):
    def test_pptx_anonymization_preserves_the_presentation(self) -> None:
        source = _pptx_bytes("DNI 12345678Z y email ana@example.com")
        engine = build_verified_rules_engine()
        analysis = engine.analyze(source)
        output, applied = engine.apply(
            source,
            analysis,
            selected_group_ids={group.id for group in analysis.groups},
            replacements={},
            remove_unsupported_content=False,
        )

        presentation = Presentation(io.BytesIO(output))
        self.assertEqual(len(applied), 2)
        self.assertEqual(
            presentation.slides[0].shapes[0].text,
            "DNI [DNI_01] y email [EMAIL_01]",
        )

    def test_pptx_template_variables_are_discovered_and_rendered(self) -> None:
        source = _pptx_bytes("Cliente: {{ cliente }}")
        analysis = analyze_template(source)
        output = render_document(source, {"cliente": "Acme"})

        presentation = Presentation(io.BytesIO(output))
        self.assertEqual(analysis.variables, ("cliente",))
        self.assertEqual(analysis.document_format, "pptx")
        self.assertEqual(presentation.slides[0].shapes[0].text, "Cliente: Acme")


if __name__ == "__main__":
    unittest.main()
