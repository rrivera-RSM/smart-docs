from __future__ import annotations

import io
import time
import unittest

from docx import Document
from fastapi.testclient import TestClient

from apps.api.main import app
from docgen.config import DOCX_MIME


def document_bytes() -> bytes:
    document = Document()
    document.add_paragraph(
        "Proyecto Horizonte. Responsable con DNI 12345678Z y correo ana@example.com."
    )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def wait_for(client: TestClient, job_id: str) -> dict[str, object]:
    for _ in range(200):
        payload = client.get(f"/api/anonymizer/{job_id}").json()
        if payload["status"] not in {"queued", "extracting", "detecting", "applying"}:
            return payload
        time.sleep(.01)
    raise AssertionError("Timeout del trabajo E2E")


class AnonymizerEndToEndTests(unittest.TestCase):
    def test_upload_review_manual_apply_verify_and_download(self) -> None:
        client = TestClient(app)
        uploaded = client.post(
            "/api/anonymizer/analyze",
            files={"file": ("piloto.docx", document_bytes(), DOCX_MIME)},
        )
        self.assertEqual(uploaded.status_code, 202)
        job = wait_for(client, uploaded.json()["job_id"])
        self.assertEqual(job["status"], "review_ready")

        block = job["blocks"][0]
        start = block["text"].index("Horizonte")
        manual = client.post(
            f"/api/anonymizer/{job['job_id']}/findings/manual",
            json={
                "block_id": block["id"],
                "start": start,
                "end": start + len("Horizonte"),
                "entity_type": "ORGANIZATION",
                "all_occurrences": True,
            },
        )
        self.assertEqual(manual.status_code, 200)
        reviewed = manual.json()

        applied = client.post(
            f"/api/anonymizer/{job['job_id']}/apply",
            json={
                "review_confirmed": True,
                "remove_unsupported_content": False,
                "decisions": [
                    {
                        "group_id": group["id"],
                        "action": "replace",
                        "replacement": group["replacement"],
                    }
                    for group in reviewed["groups"]
                ],
            },
        )
        self.assertEqual(applied.status_code, 202)
        ready = wait_for(client, job["job_id"])
        self.assertEqual(ready["status"], "ready")

        download = client.get(ready["download_url"])
        self.assertEqual(download.status_code, 200)
        rendered = Document(io.BytesIO(download.content))
        text = rendered.paragraphs[0].text
        self.assertNotIn("Horizonte", text)
        self.assertNotIn("12345678Z", text)
        self.assertNotIn("ana@example.com", text)
        self.assertIn("[EMP_01]", text)
        self.assertIn("[DNI_01]", text)
        self.assertIn("[EMAIL_01]", text)


if __name__ == "__main__":
    unittest.main()
