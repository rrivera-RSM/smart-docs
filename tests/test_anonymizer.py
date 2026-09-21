import io
import unittest

from docx import Document

from docgen.anonymizer.detectors import (
    detect_text,
    validate_dni,
    validate_nie,
    validate_spanish_iban,
)
from docgen.anonymizer.service import analyze_document, anonymize_document


def _document_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class StructuredDetectorTests(unittest.TestCase):
    def test_dni_and_nie_require_valid_control_letters(self) -> None:
        self.assertTrue(validate_dni("12345678-Z"))
        self.assertFalse(validate_dni("12345678-A"))
        self.assertTrue(validate_nie("X2482300-W"))
        self.assertFalse(validate_nie("X2482300-A"))

    def test_spanish_iban_requires_valid_checksum(self) -> None:
        self.assertTrue(validate_spanish_iban("ES91 2100 0418 4502 0005 1332"))
        self.assertFalse(validate_spanish_iban("ES00 2100 0418 4502 0005 1332"))

    def test_detector_returns_reviewable_context(self) -> None:
        findings = detect_text(
            "Contacto: ana@example.com y teléfono 612 345 678.",
            "body.paragraph[0]",
        )

        self.assertEqual(
            [finding.entity_type for finding in findings],
            ["EMAIL", "ES_PHONE"],
        )
        self.assertIn("ana@example.com", findings[0].context)


class DocumentAnonymizerTests(unittest.TestCase):
    def test_analysis_covers_body_tables_and_headers(self) -> None:
        document = Document()
        document.add_paragraph("DNI: 12345678Z")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "ana@example.com"
        document.sections[0].header.paragraphs[0].text = "Tel. 612 345 678"

        findings = analyze_document(_document_bytes(document))

        self.assertEqual(len(findings), 3)
        self.assertEqual(
            {finding.entity_type for finding in findings},
            {"ES_DNI", "EMAIL", "ES_PHONE"},
        )

    def test_selected_findings_are_replaced_without_losing_other_text(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Cliente ")
        first = paragraph.add_run("1234")
        first.bold = True
        paragraph.add_run("5678Z")
        paragraph.add_run(" permanece activo")
        source = _document_bytes(document)
        findings = analyze_document(source)

        result = anonymize_document(source, findings)
        rendered = Document(io.BytesIO(result))

        self.assertEqual(
            rendered.paragraphs[0].text,
            "Cliente [DNI_01] permanece activo",
        )
        self.assertTrue(rendered.paragraphs[0].runs[1].bold)

    def test_custom_replacement_and_metadata_cleanup(self) -> None:
        document = Document()
        document.add_paragraph("Enviar a ana@example.com")
        document.core_properties.author = "Persona sensible"
        source = _document_bytes(document)
        finding = analyze_document(source)[0]

        result = anonymize_document(
            source,
            [finding],
            {finding.id: "[CORREO_01]"},
        )
        rendered = Document(io.BytesIO(result))

        self.assertEqual(rendered.paragraphs[0].text, "Enviar a [CORREO_01]")
        self.assertEqual(rendered.core_properties.author, "")


if __name__ == "__main__":
    unittest.main()
