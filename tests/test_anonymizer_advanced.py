from __future__ import annotations

import base64
import io
import unittest
import zipfile

from docx import Document
from docx.shared import Inches

from docgen.anonymizer.detector_registry import (
    RuleBasedDetector,
    _is_explicitly_non_entity_context,
    _is_structural_label,
    validate_passport_number,
    validate_swift,
)
from docgen.anonymizer.models import AnalysisOptions, DocumentLocation
from docgen.anonymizer.resolution import canonical_value
from docgen.anonymizer.verified_engine import build_verified_rules_engine


TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _document_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class CanonicalizationTests(unittest.TestCase):
    def test_passport_variants_share_a_canonical_value(self) -> None:
        self.assertEqual(
            canonical_value("PASSPORT", "PAA-123456"),
            canonical_value("PASSPORT", "PAA 123456"),
        )

    def test_money_variants_share_a_canonical_value(self) -> None:
        self.assertEqual(
            canonical_value("MONEY", "1.234,56 €"),
            canonical_value("MONEY", "EUR 1 234,56"),
        )
        self.assertEqual(
            canonical_value("MONEY", "dos millones de euros"),
            "2E+6:eur",
        )

    def test_company_suffix_and_punctuation_are_ignored(self) -> None:
        variants = {
            canonical_value("ORGANIZATION", "Ácme Consultoría, S.L."),
            canonical_value("ORGANIZATION", "ACME CONSULTORIA SL"),
            canonical_value("ORGANIZATION", "Acme Consultoría S L"),
        }
        self.assertEqual(variants, {"acmeconsultoria"})


class DetectorPrecisionTests(unittest.TestCase):
    def test_passports_require_an_explicit_document_context(self) -> None:
        location = DocumentLocation(kind="docx", part="word/document.xml", block_index=0, label="Documento")
        findings = RuleBasedDetector().detect(
            "Pasaporte: PAA123456. Passport No. 12AB34567. "
            "Documento de viaje nº 123456789. Referencia ABC123456.",
            location,
        )
        passport_values = {
            finding.value for finding in findings if finding.entity_type == "PASSPORT"
        }

        self.assertEqual(passport_values, {"PAA123456", "12AB34567", "123456789"})
        self.assertTrue(validate_passport_number("PAA-123456"))
        self.assertFalse(validate_passport_number("PASAPORTE"))

    def test_swift_requires_a_valid_iso_country_code(self) -> None:
        location = DocumentLocation(kind="docx", part="word/document.xml", block_index=0, label="Documento")
        findings = RuleBasedDetector().detect(
            "CONTACTO TRABAJADORA CAIXESBBXXX BBVAESMMXXX",
            location,
        )
        swift_values = {finding.value for finding in findings if finding.entity_type == "SWIFT"}

        self.assertEqual(swift_values, {"CAIXESBBXXX", "BBVAESMMXXX"})
        self.assertFalse(validate_swift("CONTACTO"))
        self.assertFalse(validate_swift("TRABAJADORA"))

    def test_ner_structural_labels_are_ignored(self) -> None:
        self.assertTrue(_is_structural_label("IBAN"))
        self.assertTrue(_is_structural_label("Contacto"))
        self.assertFalse(_is_structural_label("Ana Belén García López"))

    def test_explicit_project_names_and_negative_examples_are_not_entities(self) -> None:
        text = (
            "El proyecto interno se denomina Victoria y su segunda fase se llama Amazon. "
            "Victoria y Amazon son nombres de proyecto. "
            "Negativos ambiguos como Victoria o Amazon."
        )
        for value in ("Victoria", "Amazon"):
            start = 0
            while (start := text.find(value, start)) >= 0:
                self.assertTrue(
                    _is_explicitly_non_entity_context(text, start, start + len(value)),
                    msg=f"No se filtró {value!r} en {text[start:start + 80]!r}",
                )
                start += len(value)


class AdvancedPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = build_verified_rules_engine()

    def test_passport_is_replaced_without_its_label(self) -> None:
        document = Document()
        document.add_paragraph("Número de pasaporte: PAA-123456")
        source = _document_bytes(document)
        analysis = self.engine.analyze(source)
        passport_group = next(
            group for group in analysis.groups if group.entity_type == "PASSPORT"
        )

        output, _ = self.engine.apply(
            source,
            analysis,
            selected_group_ids={passport_group.id},
            replacements={},
            remove_unsupported_content=False,
        )

        rendered = Document(io.BytesIO(output))
        self.assertEqual(rendered.paragraphs[0].text, "Número de pasaporte: [PASS_01]")

    def test_maximum_profile_groups_repeated_values(self) -> None:
        document = Document()
        document.add_paragraph(
            "Cliente: Ana García. Ana García cobrará 1.234,56 €. "
            "La transferencia de EUR 1 234,56 se hará a ana@example.com."
        )

        analysis = self.engine.analyze(_document_bytes(document))

        money_groups = [group for group in analysis.groups if group.entity_type == "MONEY"]
        self.assertEqual(len(money_groups), 1)
        self.assertEqual(len(money_groups[0].finding_ids), 2)
        self.assertEqual(money_groups[0].replacement, "[IMP_01]")
        self.assertTrue(all(group.selected_by_default for group in analysis.groups))

    def test_document_allowlist_is_exact_and_temporary(self) -> None:
        document = Document()
        document.add_paragraph("Madrid y Barcelona figuran en el contrato.")
        source = _document_bytes(document)

        filtered = self.engine.analyze(
            source,
            AnalysisOptions(allowlist=("Madrid",)),
        )
        unfiltered = self.engine.analyze(source)

        self.assertNotIn("Madrid", {finding.value for finding in filtered.findings})
        self.assertIn("Madrid", {finding.value for finding in unfiltered.findings})

    def test_manual_finding_has_priority_and_can_cover_all_occurrences(self) -> None:
        document = Document()
        document.add_paragraph("Madrid aparece aquí. Madrid aparece de nuevo.")
        analysis = self.engine.analyze(_document_bytes(document))
        block = analysis.blocks[0]

        reviewed = self.engine.add_manual_finding(
            analysis,
            block_id=block.id,
            start=0,
            end=6,
            entity_type="PERSON",
            all_occurrences=True,
        )

        manual_group = next(group for group in reviewed.groups if group.entity_type == "PERSON")
        self.assertEqual(len(manual_group.finding_ids), 2)
        self.assertEqual(manual_group.replacement, "[PERS_01]")
        self.assertIn("Revisión manual", manual_group.sources)
        self.assertIn("Gazetteer local", manual_group.sources)

    def test_repeated_form_labels_do_not_block_residual_verification(self) -> None:
        document = Document()
        document.add_paragraph("CONTACTO ana@example.com")
        document.add_paragraph("TRABAJADORA con DNI 12345678Z")
        document.add_paragraph("Campos revisados: CONTACTO, TRABAJADORA e IBAN.")
        source = _document_bytes(document)
        analysis = self.engine.analyze(source)

        self.assertFalse(
            {"CONTACTO", "TRABAJADORA", "IBAN"}
            & {finding.value for finding in analysis.findings}
        )
        output, applied = self.engine.apply(
            source,
            analysis,
            selected_group_ids={group.id for group in analysis.groups},
            replacements={},
            remove_unsupported_content=False,
        )

        self.assertGreaterEqual(len(applied), 2)
        rendered = Document(io.BytesIO(output))
        text = " ".join(paragraph.text for paragraph in rendered.paragraphs)
        self.assertIn("CONTACTO", text)
        self.assertIn("TRABAJADORA", text)
        self.assertNotIn("ana@example.com", text)
        self.assertNotIn("12345678Z", text)

    def test_images_block_download_until_explicitly_removed(self) -> None:
        document = Document()
        document.add_paragraph("DNI 12345678Z")
        document.add_picture(io.BytesIO(TINY_PNG), width=Inches(.1))
        source = _document_bytes(document)
        analysis = self.engine.analyze(source)
        selected = {group.id for group in analysis.groups}

        self.assertEqual({issue.kind for issue in analysis.coverage_issues}, {"image"})
        with self.assertRaisesRegex(ValueError, "no inspeccionables"):
            self.engine.apply(
                source,
                analysis,
                selected_group_ids=selected,
                replacements={},
                remove_unsupported_content=False,
            )

        output, applied = self.engine.apply(
            source,
            analysis,
            selected_group_ids=selected,
            replacements={},
            remove_unsupported_content=True,
        )

        self.assertEqual(len(applied), 1)
        with zipfile.ZipFile(io.BytesIO(output)) as archive:
            self.assertFalse(any(name.startswith("word/media/") for name in archive.namelist()))
        rendered = Document(io.BytesIO(output))
        self.assertEqual(rendered.paragraphs[0].text, "DNI [DNI_01]")


if __name__ == "__main__":
    unittest.main()
