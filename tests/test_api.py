import io
import time
import unittest

from docx import Document
from fastapi.testclient import TestClient

from apps.api.main import app
from docgen.config import DOCX_MIME


def _document_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _wait_for_job(client: TestClient, job_id: str) -> dict[str, object]:
    for _ in range(200):
        response = client.get(f"/api/anonymizer/{job_id}")
        if response.status_code != 200:
            raise AssertionError(response.text)
        payload = response.json()
        if payload["status"] not in {"queued", "extracting", "detecting", "applying"}:
            return payload
        time.sleep(.01)
    raise AssertionError("El trabajo no alcanzo un estado terminal.")


class SmartDocsApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.client = TestClient(app)

    def test_health_endpoint(self) -> None:
        response = self.client.get("/api/health")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")

    def test_anonymization_review_and_download_flow(self) -> None:
        source = _document_bytes(
            "DNI 12345678Z, email ana@example.com y teléfono 612 345 678."
        )
        analysis = self.client.post(
            "/api/anonymizer/analyze",
            files={"file": ("clientes.docx", source, DOCX_MIME)},
        )

        self.assertEqual(analysis.status_code, 202)
        payload = _wait_for_job(self.client, analysis.json()["job_id"])
        self.assertEqual(payload["status"], "review_ready")
        self.assertEqual(len(payload["findings"]), 3)

        decisions = [
            {
                "group_id": group["id"],
                "action": "replace",
                "replacement": group["replacement"],
            }
            for group in payload["groups"]
        ]
        applied = self.client.post(
            f"/api/anonymizer/{payload['job_id']}/apply",
            json={"decisions": decisions, "review_confirmed": True},
        )

        self.assertEqual(applied.status_code, 202)
        completed = _wait_for_job(self.client, payload["job_id"])
        self.assertEqual(completed["status"], "ready")
        self.assertEqual(completed["applied_count"], 3)
        self.assertEqual(completed["remaining_count"], 0)

        download = self.client.get(completed["download_url"])
        rendered = Document(io.BytesIO(download.content))

        self.assertEqual(download.status_code, 200)
        self.assertEqual(
            rendered.paragraphs[0].text,
            "DNI [DNI_01], email [EMAIL_01] y teléfono [TEL_01].",
        )

    def test_generator_analysis_and_download_flow(self) -> None:
        source = _document_bytes("Estimado/a {{ cliente }}")
        analysis = self.client.post(
            "/api/generator/analyze",
            files={"file": ("carta.docx", source, DOCX_MIME)},
        )

        self.assertEqual(analysis.status_code, 200)
        payload = analysis.json()
        self.assertEqual(payload["variables"], ["cliente"])

        generated = self.client.post(
            f"/api/generator/{payload['job_id']}/generate",
            json={"context": {"cliente": "Acme"}},
        )
        self.assertEqual(generated.status_code, 200)

        download = self.client.get(generated.json()["download_url"])
        rendered = Document(io.BytesIO(download.content))

        self.assertEqual(rendered.paragraphs[0].text, "Estimado/a Acme")

    def test_rejects_non_docx_uploads(self) -> None:
        response = self.client.post(
            "/api/anonymizer/analyze",
            files={"file": ("datos.txt", b"not a word file", "text/plain")},
        )

        self.assertEqual(response.status_code, 415)


if __name__ == "__main__":
    unittest.main()
