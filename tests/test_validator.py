import io
import unittest

from docx import Document

from docgen.autofix import autofix_split_placeholders
from docgen.renderer import render_docx
from docgen.template_service import analyze_template
from docgen.validator import validate_template_placeholders


def _document_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class TemplateValidationTests(unittest.TestCase):
    def test_valid_placeholder_has_no_issues(self) -> None:
        document = Document()
        document.add_paragraph("Hola {{ cliente }}")

        issues = validate_template_placeholders(_document_bytes(document))

        self.assertEqual(issues, [])

    def test_split_placeholder_is_detected_in_body(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Hola ")
        paragraph.add_run("{{ cli")
        paragraph.add_run("ente }}")

        issues = validate_template_placeholders(_document_bytes(document))

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].variable, "cliente")
        self.assertEqual(issues[0].location, "body.paragraph[0]")
        self.assertEqual(issues[0].run_texts, ("{{ cli", "ente }}"))

    def test_split_placeholders_are_detected_in_tables_and_headers(self) -> None:
        document = Document()
        cell_paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell_paragraph.add_run("{{ pro")
        cell_paragraph.add_run("yecto }}")
        header_paragraph = document.sections[0].header.paragraphs[0]
        header_paragraph.add_run("{{ cli")
        header_paragraph.add_run("ente }}")

        issues = validate_template_placeholders(_document_bytes(document))
        locations = {issue.location for issue in issues}

        self.assertEqual(len(issues), 2)
        self.assertTrue(any(".table[0]." in location for location in locations))
        self.assertTrue(any(".header." in location for location in locations))

    def test_autofix_repairs_split_placeholder_and_keeps_first_style(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        first_run = paragraph.add_run("{{ cli")
        first_run.bold = True
        paragraph.add_run("ente }}")
        source = _document_bytes(document)

        fixed, merge_count = autofix_split_placeholders(source)
        repaired = Document(io.BytesIO(fixed))

        self.assertEqual(merge_count, 1)
        self.assertEqual(validate_template_placeholders(fixed), [])
        self.assertEqual(repaired.paragraphs[0].text, "{{ cliente }}")
        self.assertTrue(repaired.paragraphs[0].runs[0].bold)


class TemplateServiceTests(unittest.TestCase):
    def test_analysis_discovers_sorted_variables(self) -> None:
        document = Document()
        document.add_paragraph("{{ proyecto }} para {{ cliente }}")

        analysis = analyze_template(_document_bytes(document))

        self.assertTrue(analysis.is_ready)
        self.assertEqual(analysis.variables, ("cliente", "proyecto"))
        self.assertEqual(analysis.issues, ())

    def test_render_replaces_template_variables(self) -> None:
        document = Document()
        document.add_paragraph("Hola, {{ cliente }}")
        source = _document_bytes(document)

        rendered = render_docx(source, {"cliente": "Acme"})
        result = Document(io.BytesIO(rendered))

        self.assertEqual(result.paragraphs[0].text, "Hola, Acme")

    def test_analysis_stops_before_extraction_when_template_needs_fix(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("{{ cli")
        paragraph.add_run("ente }}")

        analysis = analyze_template(_document_bytes(document))

        self.assertFalse(analysis.is_ready)
        self.assertEqual(analysis.variables, ())
        self.assertEqual(len(analysis.issues), 1)


if __name__ == "__main__":
    unittest.main()
