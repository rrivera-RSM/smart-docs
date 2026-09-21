import re
from dataclasses import dataclass
from typing import Any, Iterator

from docx.document import Document as DocumentType

JINJA_SIMPLE_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")


@dataclass(frozen=True)
class RunSpan:
    start: int
    end: int
    index: int
    text: str


def get_run_spans(paragraph: Any) -> list[RunSpan]:
    """Map every run to its position in the paragraph's concatenated text."""
    spans: list[RunSpan] = []
    position = 0

    for index, run in enumerate(paragraph.runs):
        text = run.text or ""
        spans.append(
            RunSpan(
                start=position,
                end=position + len(text),
                index=index,
                text=text,
            )
        )
        position += len(text)

    return spans


def find_run_index(spans: list[RunSpan], char_position: int) -> int | None:
    """Find the run that owns a character in the concatenated paragraph."""
    for span in spans:
        if span.start <= char_position < span.end:
            return span.index
    return None


def _iter_table_paragraphs(
    tables: Any,
    location_prefix: str,
) -> Iterator[tuple[str, Any]]:
    for table_index, table in enumerate(tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell_prefix = (
                    f"{location_prefix}.table[{table_index}]"
                    f".row[{row_index}].cell[{cell_index}]"
                )
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield (
                        f"{cell_prefix}.paragraph[{paragraph_index}]",
                        paragraph,
                    )
                yield from _iter_table_paragraphs(
                    cell.tables,
                    cell_prefix,
                )


def iter_document_paragraphs(
    document: DocumentType,
) -> Iterator[tuple[str, Any]]:
    """Yield paragraphs from the body, tables, headers, and footers."""
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        yield f"body.paragraph[{paragraph_index}]", paragraph

    yield from _iter_table_paragraphs(document.tables, "body")

    for section_index, section in enumerate(document.sections):
        for area_name, area in (
            ("header", section.header),
            ("footer", section.footer),
        ):
            prefix = f"section[{section_index}].{area_name}"
            for paragraph_index, paragraph in enumerate(area.paragraphs):
                yield f"{prefix}.paragraph[{paragraph_index}]", paragraph
            yield from _iter_table_paragraphs(area.tables, prefix)
