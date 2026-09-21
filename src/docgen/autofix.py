import io

from docx import Document

from docgen.placeholders import (
    JINJA_SIMPLE_RE,
    find_run_index,
    get_run_spans,
    iter_document_paragraphs,
)


def autofix_split_placeholders(docx_bytes: bytes) -> tuple[bytes, int]:
    """Merge split placeholders while preserving the first run's style."""
    document = Document(io.BytesIO(docx_bytes))
    merges = 0

    for _, paragraph in iter_document_paragraphs(document):
        # Repeat because one paragraph can contain more than one split field.
        while True:
            full_text = paragraph.text or ""
            if "{{" not in full_text or not paragraph.runs:
                break

            spans = get_run_spans(paragraph)
            changed = False

            for match in JINJA_SIMPLE_RE.finditer(full_text):
                run_start = find_run_index(spans, match.start())
                run_end = find_run_index(spans, match.end() - 1)

                if run_start is None or run_end is None or run_start == run_end:
                    continue

                runs = paragraph.runs
                merged_text = "".join(
                    runs[index].text or ""
                    for index in range(run_start, run_end + 1)
                )

                runs[run_start].text = merged_text
                for index in range(run_start + 1, run_end + 1):
                    runs[index].text = ""

                merges += 1
                changed = True
                break

            if not changed:
                break

    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), merges
