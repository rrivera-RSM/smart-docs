import io
import re
from typing import Tuple
import base64
import streamlit as st
from docx import Document

from extractor import extract_vars_docxtpl, extract_vars_fallback_regex
from renderer import render_docx
from validator import validate_template_placeholders
from pathlib import Path

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# Solo variables simples: {{ var }}
JINJA_SIMPLE_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")


def _run_spans(paragraph):
    """Devuelve spans (start, end, idx) para cada run dentro del texto
    concatenado del párrafo."""
    spans = []
    pos = 0
    for i, run in enumerate(paragraph.runs):
        t = run.text or ""
        spans.append((pos, pos + len(t), i))
        pos += len(t)
    return spans


def _find_run_index_for_char(spans, char_pos):
    for start, end, idx in spans:
        if start <= char_pos < end:
            return idx
    return None


def _iter_all_paragraphs(doc: Document):
    """
    Itera párrafos del body + tablas.
    (Headers/footers se pueden añadir más adelante si los necesitas.)
    """
    for p in doc.paragraphs:
        yield p

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def autofix_split_placeholders(docx_bytes: bytes) -> Tuple[bytes, int]:
    """
    Repara placeholders partidos en runs fusionando los runs implicados.
    Preserva el estilo del primer run (run inicial de la secuencia fusionada).

    Devuelve: (docx_bytes_reparado, numero_de_fusiones_realizadas)
    """
    doc = Document(io.BytesIO(docx_bytes))
    merges = 0

    for p in _iter_all_paragraphs(doc):
        # Repetimos hasta que ya no haya placeholders partidos en este párrafo
        while True:
            full = p.text or ""
            if "{{" not in full or not p.runs:
                break

            spans = _run_spans(p)
            changed = False

            for m in JINJA_SIMPLE_RE.finditer(full):
                start_char = m.start()
                end_char_inclusive = m.end() - 1

                run_start = _find_run_index_for_char(spans, start_char)
                run_end = _find_run_index_for_char(spans, end_char_inclusive)

                if run_start is None or run_end is None:
                    continue

                if run_start != run_end:
                    runs = p.runs

                    merged_text = "".join(
                        runs[i].text or ""
                        for i in range(run_start, run_end + 1)
                    )
                    runs[run_start].text = (
                        merged_text  # estilo del primer run se conserva
                    )

                    for i in range(run_start + 1, run_end + 1):
                        runs[i].text = ""

                    merges += 1
                    changed = True
                    break  # recomputar spans y texto

            if not changed:
                break

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), merges


st.set_page_config(
    page_title="RSM Document Generator",
    layout="centered",
    # page_icon="assets/rsm_logo.png",  # opcional
)


LOGO_PATH = Path(__file__).parent / "assets" / "rsm_logo.png"


def _img_to_base64(path: Path) -> str:
    return base64.b64encode(path.read_bytes()).decode("utf-8")


def render_header():
    b64 = _img_to_base64(LOGO_PATH) if LOGO_PATH.exists() else ""

    st.markdown(
        """
        <style>
          /* Dale aire arriba: si lo pones muy bajo puede parecer "cortado" */
          .block-container { padding-top: 2rem; }

          .rsm-header {
            display:flex;
            align-items:center;
            gap:14px;
            padding: 14px 14px;
            margin-top: 14px;
            border-radius: 12px;
            background: #ffffff;
            border: 1px solid rgba(0,21,61,0.10);
            box-shadow: 0 1px 4px rgba(0,0,0,0.04);
          }

          .rsm-logo {
            height: 48px;          /* controla altura visible */
            width: 160px;          /* evita que se “aplasten” logos muy anchos */
            object-fit: contain;   /* clave: NO recorta */
            object-position: left center;
            display:block;
          }

          .rsm-title {
            margin: 0;
            font-size: 20px;
            font-weight: 750;
            color: #00153d;
            letter-spacing: 0.2px;
          }
          .rsm-subtitle {
            margin: 2px 0 0 0;
            font-size: 12.5px;
            color: #68646c;
          }
          .rsm-accent {
            height: 3px;
            background: linear-gradient(90deg, #00153d, #009cde, #3f9c35);
            border-radius: 999px;
            margin: 10px 2px 14px 2px;
          }
        </style>
        """,
        unsafe_allow_html=True,
    )

    if b64:
        logo_html = f"<img class='rsm-logo' src='data:image/png;base64,{b64}' alt='RSM'/>"
    else:
        logo_html = "<div style='width:160px;color:#68646c;font-size:12px'>Logo no encontrado</div>"

    st.markdown(
        f"""
        <div class="rsm-header">
          {logo_html}
          <div>
            <p class="rsm-title">Document Generator</p>
            <p class="rsm-subtitle">Genera documentos a partir de plantillas .docx</p>
          </div>
        </div>
        <div class="rsm-accent"></div>
        """,
        unsafe_allow_html=True,
    )


render_header()

uploaded = st.file_uploader("Sube tu plantilla en formato Word", type=["docx"])

if uploaded:
    # Mantener en session_state el último template (útil tras auto-fix)
    if "template_bytes" not in st.session_state:
        st.session_state["template_bytes"] = uploaded.read()

    # Si el usuario sube otro archivo diferente, reseteamos el estado
    if "uploaded_name" not in st.session_state:
        st.session_state["uploaded_name"] = uploaded.name
        st.session_state["uploaded_size"] = uploaded.size
    else:
        if (
            uploaded.name != st.session_state["uploaded_name"]
            or uploaded.size != st.session_state["uploaded_size"]
        ):
            st.session_state["template_bytes"] = uploaded.read()
            st.session_state["uploaded_name"] = uploaded.name
            st.session_state["uploaded_size"] = uploaded.size

    template_bytes = st.session_state["template_bytes"]

    # 1) Validación: placeholders partidos en runs
    issues = validate_template_placeholders(template_bytes)

    if issues:
        st.error(
            "⚠️ La plantilla tiene campos partidos, lo que puede causar problemas al generar el documento."
        )

        for it in issues:
            with st.expander(f"{it.variable} en {it.location}"):
                st.write("Párrafo:", it.paragraph_preview)
                st.write("Runs implicados:", it.run_texts)

        c1, c2 = st.columns(2)

        with c1:
            if st.button("🛠️ Auto-fix "):
                fixed_bytes, merges = autofix_split_placeholders(
                    template_bytes
                )
                st.session_state["template_bytes"] = fixed_bytes

                new_issues = validate_template_placeholders(fixed_bytes)
                if new_issues:
                    st.warning(
                        f"He aplicado {merges} fusiones, quedan incidencias."
                    )
                else:
                    st.success(
                        f"Plantilla reparada ✅ (fusiones realizadas: {merges})"
                    )

        with c2:
            st.download_button(
                "⬇️ Descargar plantilla actual (reparada si aplica)",
                data=st.session_state["template_bytes"],
                file_name="plantilla_fixed.docx",
                mime=DOCX_MIME,
            )

        # Modo estricto: si todavía hay issues, no dejamos generar.
        # (Garantiza que no renderices con tags Jinja partidos en runs.)
        st.stop()

    # 2) Extraer variables (solo simples)
    try:
        variables = extract_vars_docxtpl(template_bytes)
    except Exception:
        variables = []

    if not variables:
        variables = extract_vars_fallback_regex(template_bytes)

    if not variables:
        st.warning(
            "No se encontraron variables {{ var }} simples en la plantilla."
        )
        st.stop()

    st.subheader("Campos")
    st.caption("Variables detectadas en la plantilla:")
    st.code(", ".join(variables))

    # 3) Formulario
    with st.form("fill_form"):
        context = {v: st.text_input(v) for v in variables}
        submitted = st.form_submit_button("Generar documento")

    # 4) Render + descarga
    if submitted:
        output_bytes = render_docx(template_bytes, context)
        st.success("Documento generado ✅")
        st.download_button(
            label="Descargar documento",
            data=output_bytes,
            file_name="documento_generado.docx",
            mime=DOCX_MIME,
        )
