import io
from dataclasses import dataclass

from docx import Document

from docgen.placeholders import (
    JINJA_SIMPLE_RE,
    find_run_index,
    get_run_spans,
    iter_document_paragraphs,
)


@dataclass(frozen=True)
class SplitPlaceholderIssue:
    location: str
    variable: str
    paragraph_preview: str
    run_texts: tuple[str, ...]


def _preview(text: str, max_len: int = 160) -> str:
    normalized = " ".join(text.split())
    if len(normalized) <= max_len:
        return normalized
    return f"{normalized[: max_len - 1]}…"


def validate_template_placeholders(
    docx_bytes: bytes,
) -> list[SplitPlaceholderIssue]:
    """Validate that every simple placeholder lives in a single Word run."""
    document = Document(io.BytesIO(docx_bytes))
    issues: list[SplitPlaceholderIssue] = []

    for location, paragraph in iter_document_paragraphs(document):
        full_text = paragraph.text or ""
        if "{{" not in full_text:
            continue

        spans = get_run_spans(paragraph)
        if not spans:
            continue

        for match in JINJA_SIMPLE_RE.finditer(full_text):
            run_start = find_run_index(spans, match.start())
            run_end = find_run_index(spans, match.end() - 1)

            if run_start is None or run_end is None:
                continue

            if run_start != run_end:
                run_texts = tuple(
                    span.text for span in spans[run_start : run_end + 1]
                )
                issues.append(
                    SplitPlaceholderIssue(
                        location=location,
                        variable=match.group(1),
                        paragraph_preview=_preview(full_text),
                        run_texts=run_texts,
                    )
                )

    return issues
