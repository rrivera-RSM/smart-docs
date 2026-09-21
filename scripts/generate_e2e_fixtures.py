#!/usr/bin/env python3
"""Generate deterministic DOCX and PDF fixtures for SmartDocs E2E tests."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures" / "e2e"
ANONYMIZER = FIXTURES / "anonymizer"
GENERATOR = FIXTURES / "generator"
NAVY = "17365D"
GREY = "D9D9D9"


def font_path() -> Path:
    for candidate in (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ):
        if candidate.exists():
            return candidate
    raise RuntimeError("No se ha encontrado una fuente TrueType.")


def style_run(run, size: float = 11, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Arial"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_doc(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    for name, size, bold in (
        ("Normal", 11, False),
        ("Title", 24, True),
        ("Heading 1", 16, True),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.styles["Normal"].paragraph_format.line_spacing = 1.08


def style_cell(cell, *, header: bool = False, shade: bool = False) -> None:
    properties = cell._tc.get_or_add_tcPr()
    if header or shade:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), NAVY if header else "F5F8FB")
        properties.append(shading)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), GREY)
        borders.append(element)
    properties.append(borders)
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", 110), ("start", 130), ("bottom", 110), ("end", 130)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        style_run(run, 10.5, bold=header, color="FFFFFF" if header else "000000")


def add_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.15)
    table.columns[1].width = Inches(4.65)
    table.rows[0].cells[0].text = "Dato"
    table.rows[0].cells[1].text = "Valor"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell in table.rows[0].cells:
        style_cell(cell, header=True)
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            style_cell(cell, shade=bool(index % 2))


def create_anonymizer_docx() -> Path:
    path = ANONYMIZER / "contrato_servicios.docx"
    document = Document()
    configure_doc(document)
    document.core_properties.title = "Contrato de servicios profesionales"
    document.core_properties.author = "Departamento Jurídico"
    header = document.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("Expediente de Innovación Atlántica S.L."), 9, bold=True, color=NAVY)
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("Contacto: laura.sanchez@example.com"), 8.5, color="555555")

    document.add_paragraph("Contrato de servicios profesionales", style="Title")
    document.add_paragraph(
        "Las partes formalizan este acuerdo y revisarán los datos de identificación antes de firmar."
    )
    document.add_paragraph("Partes", style="Heading 1")
    company = document.add_paragraph("Empresa: ")
    style_run(company.add_run("Innovación Atlántica S.L."), bold=True)
    company.add_run(", con NIF B12345674 y domicilio en Calle Serrano 55, 28006 Madrid.")
    person = document.add_paragraph("Cliente: ")
    style_run(person.add_run("Laura "), bold=True)
    person.add_run("Sánchez ").italic = True
    style_run(person.add_run("Pérez"), bold=True)
    person.add_run(", nacida el 14/02/1987 y de 39 años.")
    document.add_paragraph("Documento de viaje. Pasaporte: PA-123456.")

    document.add_paragraph("Datos sujetos a revisión", style="Heading 1")
    add_table(
        document,
        [
            ("DNI", "12345678Z"),
            ("Pasaporte", "PA-123456"),
            ("Correo electrónico", "laura.sanchez@example.com"),
            ("Teléfono", "+34 612 345 678"),
            ("IBAN", "ES91 2100 0418 4502 0005 1332"),
            ("Tarjeta de garantía", "4111 1111 1111 1111"),
            ("Código SWIFT", "CAIXESBBXXX"),
            ("Matrícula", "1234 BCD"),
            ("Importe", "1.234,56 €"),
        ],
    )
    document.add_paragraph("Condiciones", style="Heading 1")
    document.add_paragraph(
        "El servicio se prestará en Madrid y podrá requerir desplazamientos a Portugal. "
        "La responsable recibirá las comunicaciones en laura.sanchez@example.com."
    )
    document.save(path)
    return path


def create_generator_docx() -> Path:
    path = GENERATOR / "plantilla_carta_cliente.docx"
    document = Document()
    configure_doc(document)
    document.add_paragraph("Propuesta comercial", style="Title")
    document.add_paragraph("Plantilla breve para generar una carta a partir de datos revisados.")
    document.add_paragraph("Destinatario", style="Heading 1")
    add_table(
        document,
        [
            ("Cliente", "{{ cliente }}"),
            ("Representante", "{{ representante }}"),
            ("Importe", "{{ importe }}"),
        ],
    )
    document.add_paragraph("Resumen", style="Heading 1")
    document.add_paragraph(
        "Estimado/a {{ representante }}, adjuntamos la propuesta para "
        "{{ cliente }} por un importe total de {{ importe }}."
    )
    document.add_paragraph("Atentamente,\nEquipo SmartDocs")
    document.save(path)
    return path


def register_pdf_font() -> str:
    family = "FixtureSans"
    if family not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(family, str(font_path())))
    return family


def pdf_header(pdf: canvas.Canvas, title: str, subtitle: str) -> float:
    family = register_pdf_font()
    width, height = letter
    pdf.setFillColor(HexColor("#17365D"))
    pdf.rect(0, height - 92, width, 92, fill=1, stroke=0)
    pdf.setFillColor(HexColor("#FFFFFF"))
    pdf.setFont(family, 20)
    pdf.drawString(54, height - 52, title)
    pdf.setFont(family, 9.5)
    pdf.drawString(54, height - 72, subtitle)
    return height - 126


def pdf_rows(pdf: canvas.Canvas, rows: list[tuple[str, str]], y: float) -> None:
    family = register_pdf_font()
    for index, (label, value) in enumerate(rows):
        if index % 2:
            pdf.setFillColor(HexColor("#F2F6FA"))
            pdf.rect(48, y - 8, 516, 27, fill=1, stroke=0)
        pdf.setFillColor(HexColor("#17365D"))
        pdf.setFont(family, 9.5)
        pdf.drawString(58, y, label)
        pdf.setFillColor(HexColor("#000000"))
        pdf.setFont(family, 10)
        pdf.drawString(190, y, value)
        y -= 29


def create_digital_pdf() -> Path:
    path = ANONYMIZER / "factura_proveedor.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setTitle("Factura de proveedor")
    pdf.setAuthor("Administración")
    y = pdf_header(pdf, "Factura de proveedor", "Datos ficticios para una prueba de anonimización")
    pdf_rows(
        pdf,
        [
            ("Proveedor", "Norte Azul Consultores S.L."),
            ("NIF", "B12345674"),
            ("Representante", "Doña Ana López Martín"),
            ("Pasaporte", "PA-123456"),
            ("Correo", "ana.lopez@example.com"),
            ("Teléfono", "+34 612 345 678"),
            ("Dirección", "Calle Alcalá 85, 28009 Madrid"),
            ("IBAN", "ES91 2100 0418 4502 0005 1332"),
            ("Importe", "1.234,56 €"),
        ],
        y,
    )
    pdf.setFillColor(HexColor("#555555"))
    pdf.setFont(register_pdf_font(), 8.5)
    pdf.drawString(54, 42, "Documento de prueba. No contiene datos de una persona real.")
    pdf.save()
    return path


def create_scanned_pdf() -> Path:
    path = ANONYMIZER / "formulario_alta_escaneado.pdf"
    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path()), 64)
    body_font = ImageFont.truetype(str(font_path()), 48)
    small_font = ImageFont.truetype(str(font_path()), 34)
    draw.rectangle((0, 0, 1700, 210), fill="#17365D")
    draw.text((100, 70), "Formulario de alta", fill="white", font=title_font)
    lines = (
        "Cliente: Ana Lopez Martin",
        "DNI: 12345678Z",
        "Pasaporte: PA123456",
        "Correo: ana.lopez@example.com",
        "Telefono: 612345678",
        "Localidad: Madrid",
    )
    for index, line in enumerate(lines):
        y = 340 + index * 190
        draw.text((120, y), line, fill="black", font=body_font)
        draw.line((120, y + 82, 1560, y + 82), fill="#D9D9D9", width=3)
    draw.text(
        (120, 1900),
        "Imagen raster sin capa de texto para validar la ruta OCR local.",
        fill="#555555",
        font=small_font,
    )
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setTitle("Formulario de alta escaneado")
    pdf.drawImage(ImageReader(image), 0, 0, width=letter[0], height=letter[1])
    pdf.save()
    return path


def create_generator_pdf() -> Path:
    path = GENERATOR / "plantilla_ficha_cliente.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    y = pdf_header(pdf, "Ficha de cliente", "Plantilla PDF para el generador multiformato")
    pdf_rows(
        pdf,
        [
            ("Cliente", "{{ cliente }}"),
            ("Representante", "{{ representante }}"),
            ("Importe aprobado", "{{ importe }}"),
        ],
        y,
    )
    pdf.setFillColor(HexColor("#000000"))
    pdf.setFont(register_pdf_font(), 10)
    pdf.drawString(54, y - 130, "La ficha se completa con los valores revisados en SmartDocs.")
    pdf.save()
    return path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=("all", "docx", "pdf"), default="all")
    args = parser.parse_args()
    ANONYMIZER.mkdir(parents=True, exist_ok=True)
    GENERATOR.mkdir(parents=True, exist_ok=True)
    created: list[Path] = []
    if args.format in {"all", "docx"}:
        created.extend((create_anonymizer_docx(), create_generator_docx()))
    if args.format in {"all", "pdf"}:
        created.extend((create_digital_pdf(), create_scanned_pdf(), create_generator_pdf()))
    for path in created:
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
